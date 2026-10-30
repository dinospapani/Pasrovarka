#!/usr/bin/env python3
"""Πληρότητα φραγμάτων: βρίσκει το πιο πρόσφατο Excel του Τμήματος Αναπτύξεως Υδάτων (gov.cy),
το διαβάζει και γράφει dams_levels.json.
Χρήση:  python3 dams_build.py               (κατεβάζει από gov.cy)
        python3 dams_build.py αρχείο.xlsx   (τοπικό αρχείο, για δοκιμή)"""
import sys, re, json, io, datetime, urllib.request, urllib.parse

PAGE = 'https://www.gov.cy/moa-wdd/documents/tamieytires-neroy-fragmata-ydatodexamenes/plirotita-fragmaton/'
UA = {'User-Agent': 'Mozilla/5.0 (BiteCast dams sync; github.com/dinospapani/pasrovarka)'}
SKIP = ('ΟΛΙΚΟ', 'ΣΥΝΟΛΟ', 'ΝΟΤΙΟΣ', 'ΠΑΦΟΣ', 'ΧΡΥΣΟΧΟΥ', 'ΛΕΥΚΩΣΙΑ', 'ΦΡΑΓΜΑ', 'ΤΑΜΙΕΥΤΗΡ')


def get(url):
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=60) as r:
        return r.read(), r.geturl()


def find_xlsx():
    """Επιστρέφει λίστα με URLs .xlsx από τη σελίδα (και από υποσελίδες εγγράφων), πρώτα τα πιο πρόσφατα."""
    html, base = get(PAGE)
    html = html.decode('utf-8', 'replace')
    links = [urllib.parse.urljoin(base, h) for h in re.findall(r'href="([^"]+)"', html)]
    xl = [u for u in links if re.search(r'\.xlsx?(\?|$)', u, re.I)]
    if not xl:  # ίσως τα αρχεία είναι σε υποσελίδες
        subs = [u for u in links if 'plirotita' in u.lower() and u.rstrip('/') != PAGE.rstrip('/')][:10]
        for s in subs:
            try:
                h, b = get(s)
                xl += [urllib.parse.urljoin(b, x) for x in re.findall(r'href="([^"]+\.xlsx?[^"]*)"', h.decode('utf-8', 'replace'), re.I)]
            except Exception as e:
                print('subpage fail', s, e)
    # προτίμηση στα ελληνικά «-GR» αρχεία, όχι στα Graphs/EN
    def score(u):
        n = urllib.parse.unquote(u).upper()
        return (('GR' in n) - ('GRAPH' in n) - ('-EN' in n), n)
    seen = []
    for u in sorted(xl, key=score, reverse=True):
        if u not in seen:
            seen.append(u)
    return seen


def num(v):
    if v is None or v == '':
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(',', '.').replace('%', '').strip())
    except ValueError:
        return None


def parse(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = None
    for name in wb.sheetnames:
        if name.upper().startswith('ΙΣΤΟΣΕΛΙΔΑ-GR') or name.lower() == 'sorted':
            ws = wb[name]
            break
    if ws is None:
        ws = wb[wb.sheetnames[0]]
    date = None
    total = None
    dams = {}
    for row in ws.iter_rows(values_only=True):
        cells = list(row) + [None] * 10
        if date is None:
            for c in cells[:12]:
                if isinstance(c, datetime.datetime):
                    date = c.date().isoformat()
                    break
        # όνομα: πρώτο κείμενο στις 3 πρώτες στήλες· τα νούμερα αρχίζουν 3 στήλες μετά (Ε–M στο αρχείο του ΤΑΥ)
        ni = next((i for i in range(3) if isinstance(cells[i], str) and cells[i].strip()), None)
        if ni is None:
            continue
        nm = cells[ni]
        nm = re.sub(r'\s+', ' ', nm).strip()
        if isinstance(cells[ni + 1], (int, float)) or isinstance(cells[ni + 2], (int, float)):
            continue  # π.χ. ο πίνακας εισροής ανά μήνα κάτω από τα φράγματα
        if 'ΓΕΝΙΚΟ' in nm.upper() and num(cells[ni + 3]) and num(cells[ni + 6]) is not None:
            total = {'cap': round(num(cells[ni + 3]), 3), 'stor': round(num(cells[ni + 6]), 3)}
            total['pct'] = round(total['stor'] / total['cap'] * 100, 1)
        if not nm or any(s in nm.upper() for s in SKIP):
            continue
        cap, in24, insince, stor, pct, ly, lypct, mx, mxd = cells[ni + 3:ni + 12]
        cap, stor, pct = num(cap), num(stor), num(pct)
        if cap is None or stor is None or cap <= 0:
            continue
        if pct is None:
            pct = stor / cap * 100
        elif pct <= 1.0 and stor / cap > 0.02:  # ποσοστό ως κλάσμα
            pct *= 100
        if isinstance(mxd, datetime.datetime):
            mxd = mxd.strftime('%d/%m')
        lyp = num(lypct)
        if lyp is not None and lyp <= 1.0 and (num(ly) or 0) / cap > 0.02:
            lyp *= 100
        dams[nm] = {
            'cap': round(cap, 3), 'stor': round(stor, 3), 'pct': round(pct, 1),
            'in24': num(in24), 'inSince': num(insince),
            'ly': num(ly), 'lyPct': None if lyp is None else round(lyp, 1),
            'max': num(mx), 'maxDate': None if mxd is None else str(mxd)
        }
    if len(dams) < 8:
        raise ValueError('few dams parsed: %d' % len(dams))
    return date, dams, total


def main():
    src = None
    if len(sys.argv) > 1:
        data = open(sys.argv[1], 'rb').read()
        src = sys.argv[1].split('/')[-1]
        date, dams, total = parse(data)
    else:
        last = None
        for u in find_xlsx()[:6]:
            try:
                data, _ = get(u)
                date, dams, total = parse(data)
                src = u
                break
            except Exception as e:
                last = e
                print('skip', u, e)
        if src is None:
            raise SystemExit('no usable xlsx found (%s)' % last)
    try:
        old = json.load(open('dams_levels.json', encoding='utf-8'))
        if old.get('date') and date and old['date'] > date:
            print('existing data newer, keep', old['date'])
            return
    except Exception:
        pass
    if not total:
        c = sum(d['cap'] for d in dams.values()); st = sum(d['stor'] for d in dams.values())
        total = {'cap': round(c, 3), 'stor': round(st, 3), 'pct': round(st / c * 100, 1)}
    out = {'date': date, 'source': src,
           'total': total,
           'dams': dams}
    with open('dams_levels.json', 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, separators=(',', ':'))
    print('ok', date, len(dams), 'dams', out['total'])


if __name__ == '__main__':
    main()
